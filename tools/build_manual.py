from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "strategy_manual_v0.7.md"
OUTPUT = ROOT / "dist" / "LongVol_v0.7_Manual.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "203748"
MUTED = "555555"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
RISK_RED = "9B1C1C"
FONT = "Calibri"
CJK_FONT = FONT

REQUIRED_NEWS_MARKERS = (
    "state/news/archive/{symbol}.jsonl",
    "state/news/packets/{as_of}/{symbol}.json",
    "research_packets/{symbol}.json|txt",
    "ALPACA_NEWS_REALTIME_ENTITLED=false",
    "BENZINGA_NEWS_REALTIME_ENTITLED=false",
    "longvol news-sync",
    "longvol research-capabilities --billable-probe",
)
FORBIDDEN_NEWS_MARKERS = (
    "state/news_archive/{symbol}.json",
    "research_packets/news/{as_of}/{symbol}.json",
)


def set_run_font(run, size=None, bold=None, italic=None, color=None, name=FONT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CJK_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    for tag, attrs in (("w:tblW", {"w:w": str(total), "w:type": "dxa"}),
                       ("w:tblInd", {"w:w": "120", "w:type": "dxa"})):
        node = tbl_pr.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tbl_pr.append(node)
        for key, value in attrs.items():
            node.set(qn(key), value)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, 9, color=MUTED)
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts"); fonts.set(qn("w:ascii"), FONT); fonts.set(qn("w:hAnsi"), FONT); fonts.set(qn("w:eastAsia"), CJK_FONT)
    rpr.extend([fonts, color, underline])
    text_node = OxmlElement("w:t"); text_node.text = text
    run.extend([rpr, text_node]); link.append(run); paragraph._p.append(link)


def add_inline(paragraph, text):
    token_re = re.compile(r"\[([^\]]+)\]\((https?://[^)]+)\)|(https?://[^\s，。；、)]+)")
    pos = 0
    for match in token_re.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()].replace("`", ""))
            set_run_font(run)
        label = match.group(1) or match.group(3)
        url = match.group(2) or match.group(3)
        add_hyperlink(paragraph, label, url)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:].replace("`", ""))
        set_run_font(run)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT; normal.font.size = Pt(11); normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.paragraph_format.space_before = Pt(0); normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in tokens.items():
        style = doc.styles[name]
        style.font.name = FONT; style.font.size = Pt(size); style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = FONT; style.font.size = Pt(11)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style.paragraph_format.left_indent = Inches(.375)
        style.paragraph_format.first_line_indent = Inches(-.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25
        style.paragraph_format.keep_together = True
        style.paragraph_format.keep_with_next = False

    # Use a plain paragraph style with a literal bullet.  LibreOffice's headless
    # renderer can occasionally lay out consecutive paragraphs based on Word's
    # built-in List Bullet style on the same baseline.
    manual_bullet = doc.styles.add_style("Manual Bullet", WD_STYLE_TYPE.PARAGRAPH)
    manual_bullet.base_style = normal
    manual_bullet.font.name = FONT; manual_bullet.font.size = Pt(11)
    manual_bullet._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    manual_bullet.paragraph_format.left_indent = Inches(.375)
    manual_bullet.paragraph_format.first_line_indent = Inches(-.188)
    manual_bullet.paragraph_format.space_after = Pt(4)
    manual_bullet.paragraph_format.line_spacing = 1.25
    manual_bullet.paragraph_format.keep_together = True

    manual_number = doc.styles.add_style("Manual Number", WD_STYLE_TYPE.PARAGRAPH)
    manual_number.font.name = FONT; manual_number.font.size = Pt(10.5)
    manual_number._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    manual_number.paragraph_format.left_indent = Inches(.30)
    manual_number.paragraph_format.first_line_indent = Inches(-.30)
    manual_number.paragraph_format.space_after = Pt(2)
    manual_number.paragraph_format.line_spacing = 1.15

    if "Code Block" not in [s.name for s in doc.styles]:
        code = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = doc.styles["Code Block"]
    code.font.name = "DejaVu Sans Mono"; code.font.size = Pt(8.5)
    code._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    code.paragraph_format.left_indent = Inches(.15); code.paragraph_format.right_indent = Inches(.1)
    code.paragraph_format.space_before = Pt(4); code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.keep_together = True

    callout = doc.styles.add_style("Lead Callout", WD_STYLE_TYPE.PARAGRAPH)
    callout.font.name = FONT; callout.font.size = Pt(10.5); callout.font.bold = True
    callout.font.color.rgb = RGBColor.from_string(NAVY)
    callout._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    callout.paragraph_format.left_indent = Inches(.18); callout.paragraph_format.right_indent = Inches(.18)
    callout.paragraph_format.space_before = Pt(8); callout.paragraph_format.space_after = Pt(10)
    callout.paragraph_format.line_spacing = 1.25


def add_shading(paragraph, fill=CALLOUT, border=BLUE):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); ppr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left"); left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8"); left.set(qn("w:color"), border)
    borders.append(left); ppr.append(borders)


def add_cover(doc):
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.text = "LONGVOL RESEARCH SYSTEM"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        set_run_font(run, 8.5, bold=True, color=MUTED)
    add_page_field(section.footer.paragraphs[0])

    spacer = doc.add_paragraph(); spacer.paragraph_format.space_after = Pt(108)
    kicker = doc.add_paragraph(); kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    set_run_font(kicker.add_run("RESEARCH · RISK CONTROL · LOCAL DEPLOYMENT"), 10, bold=True, color=BLUE)
    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    set_run_font(title.add_run("LongVol v0.7"), 30, bold=True, color=NAVY)
    subtitle = doc.add_paragraph(); subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(6)
    set_run_font(subtitle.add_run("US Panic-Dislocation Rebound Strategy"), 17, bold=True, color=DARK_BLUE)
    sub2 = doc.add_paragraph(); sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub2.paragraph_format.space_after = Pt(54)
    set_run_font(sub2.add_run("Production Deployment, Rules, and Research Manual"), 12, color=MUTED)
    for label, value in (("Application version", "0.7.0"), ("Strategy rules", "0.7.0"),
                         ("Date", "2026-09-11"),
                         ("Scope", "Research, risk, paper/live broker workflow"),
                         ("Execution", "Paper by default; live requires explicit approvals")):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        set_run_font(p.add_run(f"{label}  "), 10, bold=True, color=NAVY)
        set_run_font(p.add_run(value), 10, color=MUTED)
    warning = doc.add_paragraph(); warning.alignment = WD_ALIGN_PARAGRAPH.CENTER
    warning.paragraph_format.space_before = Pt(30)
    set_run_font(warning.add_run("Research system only. Options can lose the entire premium."), 9.5, italic=True, color=RISK_RED)
    doc.add_page_break()


def add_table(doc, rows):
    cols = len(rows[0])
    if cols == 2:
        widths = [2700, 6660]
    elif cols == 3:
        widths = ([900, 5700, 2760]
                  if rows[0] == ["Priority", "Trigger", "Default action"]
                  else [1800, 2500, 5060])
    elif cols == 4:
        widths = [2200, 2200, 3200, 1760]
    else:
        widths = [9360 // cols] * cols
        widths[-1] += 9360 - sum(widths)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for r_idx, values in enumerate(rows):
        for c_idx, value in enumerate(values):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_idx == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 or len(value) < 12 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(value.replace("`", ""))
            set_run_font(run, 9.2, bold=(r_idx == 0), color=NAVY if r_idx == 0 else None)
    after = doc.add_paragraph(); after.paragraph_format.space_after = Pt(2)


def parse_body(doc, source):
    lines = source.splitlines()
    start = next(i for i, line in enumerate(lines)
                 if line.startswith("# 1.") or line.startswith("## 1."))
    lines = lines[start:]
    i = 0
    in_code = False
    code_lines = []
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        if line.startswith("```"):
            if in_code:
                p = doc.add_paragraph(style="Code Block")
                add_shading(p, LIGHT_GRAY, "D7DBE2")
                run = p.add_run("\n".join(code_lines)); set_run_font(run, 8.5, name="DejaVu Sans Mono")
                code_lines = []; in_code = False
            else:
                in_code = True
            i += 1; continue
        if in_code:
            code_lines.append(raw); i += 1; continue
        if not line:
            i += 1; continue
        if line == "[PAGE BREAK]":
            doc.add_page_break()
            i += 1; continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1]):
            table_rows = [[x.strip() for x in line.strip("|").split("|")]]
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                table_rows.append([x.strip() for x in lines[i].strip("|").split("|")])
                i += 1
            add_table(doc, table_rows)
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 3"); add_inline(p, line[4:])
        elif line.startswith("## "):
            p = doc.add_paragraph(style="Heading 2"); add_inline(p, line[3:])
        elif line.startswith("# "):
            p = doc.add_paragraph(style="Heading 1"); add_inline(p, line[2:])
        elif line.startswith("> "):
            p = doc.add_paragraph(style="Lead Callout"); add_shading(p); add_inline(p, line[2:])
        elif line.startswith("- "):
            p = doc.add_paragraph(style="Manual Bullet")
            add_inline(p, "• " + line[2:])
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="Manual Number"); add_inline(p, line)
        else:
            p = doc.add_paragraph(); add_inline(p, line)
        i += 1


def audit(doc, source):
    section = doc.sections[0]
    assert round(section.left_margin.inches, 3) == 1.0
    assert round(section.page_width.inches, 3) == 8.5
    for table in doc.tables:
        tbl_w = table._tbl.tblPr.find(qn("w:tblW"))
        assert tbl_w is not None and int(tbl_w.get(qn("w:w"))) == 9360
        assert table._tbl.tblPr.find(qn("w:tblInd")) is not None
    assert not any(p.text.startswith("- ") for p in doc.paragraphs)
    for marker in REQUIRED_NEWS_MARKERS:
        assert marker in source, f"manual is missing news contract marker: {marker}"
    for marker in FORBIDDEN_NEWS_MARKERS:
        assert marker not in source, f"manual still contains stale news path: {marker}"


def main():
    doc = Document()
    style_document(doc)
    add_cover(doc)
    source = SOURCE.read_text(encoding="utf-8")
    lead = next(line[2:] for line in source.splitlines() if line.startswith("> "))
    p = doc.add_paragraph(style="Lead Callout"); add_shading(p); add_inline(p, lead)
    parse_body(doc, source)
    for section in doc.sections:
        header = section.header.paragraphs[0]
        if not header.text:
            header.text = "LONGVOL v0.7  |  STRATEGY AND DEPLOYMENT MANUAL"
        for run in header.runs:
            set_run_font(run, 8.5, bold=True, color=MUTED)
        if not section.footer.paragraphs[0].text:
            add_page_field(section.footer.paragraphs[0])
    audit(doc, source)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "LongVol v0.7 Panic-Dislocation Rebound Strategy"
    doc.core_properties.subject = "Production deployment, strategy rules, and research manual"
    doc.core_properties.author = "LongVol Research"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
